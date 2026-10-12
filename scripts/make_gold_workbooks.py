"""
CPRT - build BLINDED gold-labeling workbooks.
  gold100_rater1.xlsx  : 100 cases for you (Rater 1)
  gold30_rater2.xlsx   : 30 of those same cases for Rater 2 (separate file = blinded to you)
  gold100_KEY.csv      : silver labels  -> DO NOT OPEN until both raters finish
Run from project root:  python3 cprt/scripts/make_gold_workbooks.py
"""
import os, numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment

SEED, N_NEG, N_POS, R2_NEG, R2_POS = 42, 35, 65, 10, 20
EXCLUDE = "gold_sample_v1_defA.xlsx"            # skip cases already labeled
OUT = "gold_labeling"; os.makedirs(OUT, exist_ok=True)
rng = np.random.default_rng(SEED)

ops = pd.read_csv("operations_5k.csv.gz")
sil = pd.read_csv("silver_labels_v3.csv.gz")
dx  = pd.read_csv("diagnosis_5k.csv.gz")
med = pd.read_csv("medications_cardiopulm_5k.csv.gz")
lab = pd.read_csv("labs_cardiopulm_5k.csv.gz")

if os.path.exists(EXCLUDE):
    sil = sil[~sil.op_id.isin(pd.read_excel(EXCLUDE, engine="openpyxl").op_id)]
sil["rules_fired"] = sil["rules_fired"].fillna("none")

# ---- stratified sample: 35 negatives + 65 positives spread across rule patterns ----
neg = sil[sil.silver_label == 0].sample(N_NEG, random_state=SEED)
pos_pool = sil[sil.silver_label == 1]
share = pos_pool.rules_fired.value_counts(normalize=True)
parts = [pos_pool[pos_pool.rules_fired == k].sample(min(max(5, round(N_POS * s)), (pos_pool.rules_fired == k).sum()), random_state=SEED)
         for k, s in share.items()]
pos = pd.concat(parts).sample(frac=1, random_state=SEED).head(N_POS)
samp = pd.concat([neg, pos]).sample(frac=1, random_state=SEED).reset_index(drop=True)   # shuffle order

# ---- clinical summary per case (NO silver info) ----
icd_col = next(c for c in dx.columns if "icd" in c.lower())
dxs = dx.assign(c=dx[icd_col].astype(str).str[:3]).groupby("subject_id")["c"].apply(lambda s: "; ".join(sorted(set(s))))
drs = med.groupby("subject_id")["drug_name"].apply(lambda s: "; ".join(sorted(set(s.astype(str).str.lower()))))
lab["value"] = pd.to_numeric(lab["value"], errors="coerce")
lbs = lab.groupby(["subject_id", "item_name"])["value"].max().reset_index() \
         .groupby("subject_id").apply(lambda d: "; ".join(f"{i}={v:.2f}" for i, v in zip(d.item_name, d.value)))

keep = [c for c in ["op_id", "subject_id", "age", "sex", "asa", "emop", "department", "icd10_pcs"] if c in ops]
case = samp[["op_id"]].merge(ops[keep], on="op_id", how="left")
case["diagnoses_3char"] = case.subject_id.map(dxs).fillna("")
case["drugs"] = case.subject_id.map(drs).fillna("")
case["labs_max"] = case.subject_id.map(lbs).fillna("")
case.insert(0, "case_no", range(1, len(case) + 1))

GUIDE = [
 ("QUESTION", "Should this patient ENTER the 2024 ACC/AHA perioperative cardiovascular flowchart (or has pulmonary disease per CPRT)?"),
 ("1 = YES", "CV risk factor (incl. HTN), CV disease, CV symptom (dyspnea R06, chest pain R07, syncope R55), or chronic pulmonary disease / lung cancer surgery"),
 ("0 = NO", "None of the above. Enoxaparin/heparin alone (prophylaxis) or aspirin alone (primary prevention) = 0"),
 ("U = UNSURE", "Use when truly uncertain; write why in notes"),
 ("Confidence", "High / Medium / Low"),
 ("Rules", "Label ALONE. Do not discuss cases with the other rater until both are done."),
]

def write(df, path, rater):
    wb = Workbook(); g = wb.active; g.title = "Guide"
    for r, (k, v) in enumerate(GUIDE, 1):
        g.cell(r, 1, k).font = Font(bold=True); g.cell(r, 2, v).alignment = Alignment(wrap_text=True)
    g.column_dimensions["A"].width, g.column_dimensions["B"].width = 14, 110
    ws = wb.create_sheet("Cases")
    cols = list(df.columns) + [f"gold_{rater}", f"confidence_{rater}", f"notes_{rater}"]
    ws.append(cols)
    for row in df.itertuples(index=False): ws.append(list(row) + ["", "", ""])
    fill = PatternFill("solid", fgColor="FFF2CC")
    for c in range(1, len(cols) + 1):
        ws.cell(1, c).font = Font(bold=True)
        if c > len(df.columns):
            for r in range(1, len(df) + 2): ws.cell(r, c).fill = fill
    n = len(df) + 1; L = lambda i: ws.cell(1, i).column_letter
    for i, opts in [(len(cols) - 2, '"1,0,U"'), (len(cols) - 1, '"High,Medium,Low"')]:
        dv = DataValidation(type="list", formula1=opts, allow_blank=True); ws.add_data_validation(dv)
        dv.add(f"{L(i)}2:{L(i)}{n}")
    for i, w in zip(range(1, len(cols) + 1), [8, 11, 11, 5, 4, 4, 5, 7, 9, 40, 50, 50, 9, 11, 40]):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = "B2"; wb.save(path)

write(case, f"{OUT}/gold100_rater1.xlsx", "r1")
r2_ids = pd.concat([neg.sample(R2_NEG, random_state=SEED), pos.sample(R2_POS, random_state=SEED)]).op_id
write(case[case.op_id.isin(r2_ids)].drop(columns="case_no").reset_index(drop=True)
      .rename_axis(None).pipe(lambda d: d.insert(0, "case_no", range(1, len(d) + 1)) or d),
      f"{OUT}/gold30_rater2.xlsx", "r2")
samp[["op_id", "silver_label", "rules_fired"]].assign(in_rater2=samp.op_id.isin(r2_ids).astype(int)) \
    .to_csv(f"{OUT}/gold100_KEY.csv", index=False)
print(f"Saved in {OUT}/ : gold100_rater1.xlsx (n={len(case)}), gold30_rater2.xlsx (n={len(r2_ids)}), gold100_KEY.csv")
print("Silver mix hidden in KEY:", samp.silver_label.value_counts().to_dict())
