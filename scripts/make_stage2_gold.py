"""
CPRT Stage 2 gold check - BLINDED triage workbooks (Jev AI track hidden).
  gold_labeling/stage2_gold_rater1.xlsx  (all cases)
  gold_labeling/stage2_gold_rater2.xlsx  (same cases, separate file -> kappa)
  gold_labeling/stage2_gold_KEY.csv      (Jev AI answers - DO NOT OPEN until both raters finish)
Run from project root:  python3 cprt/scripts/make_stage2_gold.py
"""
import os, argparse, numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ap = argparse.ArgumentParser()
ap.add_argument("--seed", type=int, default=42)
ap.add_argument("--tag", default="stage2_gold", help="file prefix, e.g. stage2_valid")
ap.add_argument("--exclude", default="", help="CSV with op_id column: cases already used (development set)")
ap.add_argument("--raters", type=int, default=2)
A = ap.parse_args()
SEED, DAY, OUT = A.seed, 1440, "gold_labeling"
# stratum -> (filter on Jev AI output, n)
PLAN = [
    ("RED_acs_troponin",   lambda t: t.track.eq("RED") & t.red_flags.str.contains("ACS"), 6),
    ("RED_stroke",         lambda t: t.track.eq("RED") & t.red_flags.str.contains("stroke"), 6),
    ("YELLOW_biomarker",   lambda t: t.track.eq("YELLOW") & t.final_rule.eq("B0"), 5),
    ("YELLOW_pulmonary",   lambda t: t.track.eq("YELLOW") & t.pulmonary_flag & ~t.final_rule.eq("B0"), 5),
    ("YELLOW_stroke_rel",  lambda t: t.audit.str.contains("R4s:"), 3),
    ("YELLOW_arrhythmia",  lambda t: t.audit.str.contains("Y0:"), 3),
    ("GREEN_low_rcri",     lambda t: t.track.eq("GREEN") & t.final_rule.eq("E0"), 5),
    ("GREEN_noCV_rules",   lambda t: t.track.eq("GREEN") & t.final_rule.eq("N1") & t.stage1_route.eq("RULES"), 6),
    ("GREEN_noCV_model",   lambda t: t.track.eq("GREEN") & t.stage1_route.eq("MODEL_SAFETY_NET"), 3),
    ("GREEN_emergency",    lambda t: t.priority.str.startswith("EMERGENCY"), 3),
    ("OUT_OF_SCOPE",       lambda t: t.track.eq("OUT_OF_SCOPE"), 2),
    ("ENRICH_HF_or_AF",    lambda t: t.hf_af, 5),   # I50/I48 not RED -> are we missing decompensation/instability?
]

t = pd.read_csv("outputs/stage2_jev_triage.csv").fillna({"audit": "", "red_flags": "", "priority": ""})
t["final_rule"] = t.audit.str.split(r" \|\| ").str[-1].str.split(":").str[0]
ops = pd.read_csv("operations_5k.csv.gz")
dx = pd.read_csv("diagnosis_5k.csv.gz"); icd = next(c for c in dx.columns if "icd" in c.lower())
dx["code"] = dx[icd].astype(str).str.replace(".", "", regex=False).str[:3].str.upper()
med = pd.read_csv("medications_cardiopulm_5k.csv.gz"); lab = pd.read_csv("labs_cardiopulm_5k.csv.gz")
lab["value"] = pd.to_numeric(lab["value"], errors="coerce")
t = t.merge(ops[["op_id", "orin_time"]], on="op_id")
hfaf = dx.merge(t[["op_id", "subject_id", "orin_time"]], on="subject_id")
hfaf = hfaf[(hfaf.code.isin(["I50", "I48"])) & (hfaf.chart_time <= hfaf.orin_time)].op_id.unique()
t["hf_af"] = t.op_id.isin(hfaf) & ~t.track.isin(["RED", "OUT_OF_SCOPE"])

taken = set(pd.read_csv(A.exclude).op_id) if A.exclude else set()
print(f"excluding {len(taken)} development-set cases")
rows = []
for name, f, n in PLAN:
    pool = t[f(t) & ~t.op_id.isin(taken)]
    pick = pool.sample(min(n, len(pool)), random_state=SEED)
    taken |= set(pick.op_id); rows.append(pick.assign(stratum=name))
    print(f"{name:20s} available={len(pool):5d} sampled={len(pick)}")
S = pd.concat(rows).sample(frac=1, random_state=SEED).reset_index(drop=True)

def summary(o):
    t0 = o.orin_time
    d = dx[(dx.subject_id == o.subject_id) & (dx.chart_time <= t0)]
    d = d.assign(days=((t0 - d.chart_time) / DAY).round().astype(int)).sort_values("days").drop_duplicates("code")
    dx_s = "; ".join(f"{c}({x}d)" if x <= 90 else c for c, x in zip(d.code, d.days))
    m = med[(med.subject_id == o.subject_id) & (med.chart_time <= t0) & (med.chart_time >= t0 - 30 * DAY)]
    rc = next((c for c in med.columns if "route" in c.lower()), None)
    m = m.assign(days=((t0 - m.chart_time) / DAY).round().astype(int)).sort_values("days")
    m = m.drop_duplicates(["drug_name"] + ([rc] if rc else []))
    md_s = "; ".join(f"{dn}{' ['+str(r)+']' if rc else ''}({x}d)" for dn, r, x in zip(m.drug_name, m[rc] if rc else [""] * len(m), m.days))
    l = lab[(lab.subject_id == o.subject_id) & (lab.chart_time <= t0) & (lab.chart_time >= t0 - 7 * DAY)]
    lb_s = "; ".join(f"{k}: max {v.max():.2f} / last {v.iloc[-1]:.2f}" for k, v in l.sort_values("chart_time").groupby("item_name").value
                     if k in ("troponin_i", "troponin_t", "ckmb", "creatinine", "potassium", "pao2", "ph"))
    return dx_s, md_s, lb_s

info = ops.set_index("op_id")
cases = []
for i, o in enumerate(S.itertuples(), 1):
    r = info.loc[o.op_id]; dx_s, md_s, lb_s = summary(o)
    cases.append(dict(case_no=i, op_id=o.op_id, age=r.get("age"), sex=r.get("sex"), asa=r.get("asa"),
                      emergency=int(r.get("emop", 0) == 1), department=r.get("department"), procedure_pcs=r.get("icd10_pcs"),
                      preop_diagnoses=dx_s, preop_drugs_30d=md_s, preop_labs_7d=lb_s))
C = pd.DataFrame(cases)

GUIDE = [
 ("TASK", "Apply the 2024 ACC/AHA perioperative flowchart to each case as YOU would in pre-op clinic. Use only the data shown."),
 ("Codes", "Diagnosis 'I63(12d)' = coded 12 days before surgery (shown only if <=90 d). Drugs: [route](days before surgery). Labs: last 7 days."),
 ("RED", "Defer / urgent cardiology + MDT: ACS, unstable arrhythmia, decompensated HF, recent stroke (<3 mo) before elective surgery"),
 ("YELLOW", "Proceed only after an action: test (ECG/echo/BNP), consult (cardiology/pulm/ACHD/CIED/neuro), or MDT"),
 ("GREEN", "Proceed to surgery with standard care (includes emergencies: proceed)"),
 ("OUT_OF_SCOPE", "Cardiac surgery - ACC/AHA noncardiac guideline does not apply"),
 ("destination", "Pick the single most important next step"),
 ("Low-risk policy", "Low-risk procedure (eye/cataract, skin/breast incl. lumpectomy, thyroid/parathyroid, dental, endoscopy, minor gyn/OB, superficial, minor ear) with recent MI/ACS or stroke -> may PROCEED (GREEN, under MAC/local/topical). Not for cardiac arrest, unstable arrhythmia or decompensated HF."),
 ("Stroke timing", "Stroke/hemorrhage <=14 d -> RED. 15-90 d before cancer surgery -> YELLOW (MDT timing). 15-90 d before elective unrelated surgery -> RED (delay >=3 mo)."),
 ("No CV finding", "No CV risk factor/disease/symptom -> GREEN even if the surgery is major (ACC/AHA Step 1)."),
 ("Breast cancer", "Possible breast malignancy should not be postponed long: prefer YELLOW (work-up) over RED unless an absolute contraindication exists."),
 ("Rules", "Label ALONE. Do not discuss with the other rater. Do not open the KEY file."),
]
DEST = '"Proceed,ECG,Echo (TTE),Obtain BNP/troponin,Cardiology consult,Pulmonary optimization,Neuro/vascular team,ACHD/CIED/other consult,MDT timing decision,Defer surgery,Cardiac surgery pathway"'

def write(path, r):
    wb = Workbook(); g = wb.active; g.title = "Guide"
    for k, (a, b) in enumerate(GUIDE, 1):
        g.cell(k, 1, a).font = Font(name="Arial", bold=True); g.cell(k, 2, b).font = Font(name="Arial")
        g.cell(k, 2).alignment = Alignment(wrap_text=True)
    g.cell(len(GUIDE) + 2, 1, "Example").font = Font(name="Arial", bold=True)
    g.cell(len(GUIDE) + 2, 2, "track=YELLOW | destination=Echo (TTE) | confidence=High | notes='I35 aortic valve disease, no echo in record'").font = Font(name="Arial")
    g.column_dimensions["A"].width, g.column_dimensions["B"].width = 14, 120
    ws = wb.create_sheet("Cases"); cols = list(C.columns) + [f"track_{r}", f"destination_{r}", f"confidence_{r}", f"notes_{r}"]
    ws.append(cols)
    for row in C.itertuples(index=False): ws.append(list(row) + ["", "", "", ""])
    y = PatternFill("solid", fgColor="FFF2CC")
    for c in range(1, len(cols) + 1):
        ws.cell(1, c).font = Font(name="Arial", bold=True)
        for rr in range(2, len(C) + 2):
            ws.cell(rr, c).font = Font(name="Arial"); ws.cell(rr, c).alignment = Alignment(wrap_text=True, vertical="top")
            if c > len(C.columns): ws.cell(rr, c).fill = y
        if c > len(C.columns): ws.cell(1, c).fill = y
    L = lambda i: ws.cell(1, i).column_letter; n = len(C) + 1; b = len(C.columns)
    for i, opts in [(b + 1, '"RED,YELLOW,GREEN,OUT_OF_SCOPE"'), (b + 2, DEST), (b + 3, '"High,Medium,Low"')]:
        dv = DataValidation(type="list", formula1=opts, allow_blank=True); ws.add_data_validation(dv); dv.add(f"{L(i)}2:{L(i)}{n}")
    for i, w in enumerate([6, 11, 5, 4, 4, 5, 6, 9, 45, 50, 38, 13, 22, 11, 35], 1): ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = "C2"; wb.save(path)

os.makedirs(OUT, exist_ok=True)
for i in range(1, A.raters + 1): write(f"{OUT}/{A.tag}_rater{i}.xlsx", f"r{i}")
S[["op_id", "stratum", "track", "priority", "destination", "red_flags", "modifiers", "final_rule", "audit"]] \
    .to_csv(f"{OUT}/{A.tag}_KEY.csv", index=False)
print(f"\nSaved {len(C)} cases -> {OUT}/{A.tag}_rater1..{A.raters}.xlsx + {A.tag}_KEY.csv")
print("Jev AI track mix (hidden in KEY):", S.track.value_counts().to_dict())
